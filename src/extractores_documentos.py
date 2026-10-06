"""
Módulo de extracción de contenido a partir de documentos estructurados (PowerPoint .pptx, etc.)
para el pipeline académico de LectureFlow.
"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Union, List

logger = logging.getLogger("extractores_documentos")


def _extraer_texto_de_forma(shape) -> List[str]:
    """Extrae texto de una forma, manejando marcos de texto, tablas y formas agrupadas."""
    lineas = []

    # 1. Marcos de texto regulares
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            texto_p = "".join(run.text for run in p.runs).strip() or p.text.strip()
            if texto_p:
                lineas.append(texto_p)

    # 2. Tablas embebidas en la diapositiva
    elif shape.has_table:
        for row in shape.table.rows:
            celdas = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if celdas:
                lineas.append(" | ".join(celdas))

    # 3. Formas agrupadas (GroupShape)
    elif shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
        try:
            for sub_shape in shape.shapes:
                lineas.extend(_extraer_texto_de_forma(sub_shape))
        except Exception:
            pass

    return lineas


def extraer_texto_pptx(ruta_archivo: Union[Path, str]) -> str:
    """Extrae el contenido textual y notas de orador de una presentación PowerPoint (.pptx).

    Organiza el contenido agrupado por diapositiva ('### Diapositiva N') con marcas
    temporales estimadas para integrarse sin fricción con el sintetizador de apuntes.

    Args:
        ruta_archivo: Ruta al archivo .pptx

    Returns:
        Cadena con el texto estructurado en Markdown.
    """
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ImportError(
            "La librería 'python-pptx' no está instalada. "
            "Ejecuta: venv/bin/pip install python-pptx"
        ) from exc

    ruta = Path(ruta_archivo).resolve()
    if not ruta.exists():
        raise FileNotFoundError(f"No se encuentra la presentación en: {ruta}")

    logger.info("Iniciando extracción de presentación: %s", ruta.name)
    prs = Presentation(str(ruta))

    bloques_diapositivas: List[str] = []

    for idx, slide in enumerate(prs.slides, start=1):
        minutos = (idx - 1) * 2
        timestamp_estimado = f"[{minutos // 60:02d}:{minutos % 60:02d}:00]"

        lineas_slide: List[str] = []
        lineas_slide.append(f"### Diapositiva {idx}")
        lineas_slide.append(f"*Marca temporal de referencia: {timestamp_estimado}*")

        # Extraer texto de todas las formas
        textos_cuerpo: List[str] = []
        for shape in slide.shapes:
            textos_cuerpo.extend(_extraer_texto_de_forma(shape))

        # Eliminar líneas duplicadas consecutivas manteniendo el orden
        lineas_unicas: List[str] = []
        for t in textos_cuerpo:
            if not lineas_unicas or lineas_unicas[-1] != t:
                lineas_unicas.append(t)

        if lineas_unicas:
            lineas_slide.extend([f"- {t}" if not t.startswith(("-", "*", "#")) else t for t in lineas_unicas])
        else:
            lineas_slide.append("(Sin contenido textual detectable en la diapositiva)")

        # Extraer notas del orador si existen
        if slide.has_notes_slide:
            notes_slide = slide.notes_slide
            tf = notes_slide.notes_text_frame
            if tf and tf.text.strip():
                notas_limpias = tf.text.strip()
                lineas_slide.append("\n> 📝 **Notas del orador:**")
                for n_line in notas_limpias.splitlines():
                    if n_line.strip():
                        lineas_slide.append(f"> {n_line.strip()}")

        bloques_diapositivas.append("\n".join(lineas_slide))

    contenido_final = "\n\n---\n\n".join(bloques_diapositivas)
    logger.info(
        "Extracción completada: %d diapositivas procesadas (%d caracteres).",
        len(bloques_diapositivas),
        len(contenido_final),
    )
    return contenido_final
